"""Bước 4 của Prompt Visuals dùng ĐÚNG bảng cảnh của tab Tự động.

Chủ dự án, 21/09/2026: file Excel *"không thể edit được vì đâu biết scene nào ở
giây nào"*, và cả tab *"không đạt hiệu quả như việc sản xuất video tự động"*.
Bốn thứ phải đúng, bài nào cũng gọi thẳng hàm mà cái nút gọi:

1. cột **Giây** (`01:12–01:19 · 7s`) hiện ngay trên bảng, thiếu mốc thì "—";
2. **Mở tệp cũ** nạp lại được file Excel hôm trước (kiểm cột theo TÊN);
3. **tick vài cảnh → tạo lại đúng mấy cảnh ấy**, giá nói trước khi gửi;
4. **Lưu chỉnh sửa** vẫn ghi theo tên cột, không đoán vị trí.

Không bài nào gọi mạng: `start_batch` bị thay bằng một cái sổ.
"""

from __future__ import annotations

import importlib.util
import os
from types import SimpleNamespace

import pytest

pytest.importorskip("PyQt5.QtWidgets", reason="máy chạy test không có giao diện")
pytest.importorskip("openpyxl")

GOC = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


# ── Cột Giây: phần nghĩ, không cần cửa sổ ────────────────────────────────────

class TestChuGiay:
    def test_co_moc_thi_hien_tu_giay_nao_toi_giay_nao(self):
        from ui_qt.bang_canh_auto import chu_giay

        assert chu_giay({"srt_start": "00:01:12,000",
                         "srt_end": "00:01:19,000",
                         "duration": 7.0}) == "01:12–01:19 · 7s"

    def test_moc_dang_mm_ss_cua_file_cu_van_doc_duoc(self):
        from ui_qt.bang_canh_auto import chu_giay

        assert chu_giay({"srt_start": "00:00", "srt_end": "00:08"}) \
            == "00:00–00:08 · 8s"

    def test_phim_dai_hon_mot_tieng_thi_co_gio(self):
        from ui_qt.bang_canh_auto import chu_giay

        assert chu_giay({"srt_start": "01:02:03,000",
                         "srt_end": "01:02:11,000"}) == "1:02:03–1:02:11 · 8s"

    def test_khong_co_moc_thi_gach_ngang_chu_khong_bia_so(self):
        from ui_qt.bang_canh_auto import chu_giay

        assert chu_giay({"scene_id": 5}) == "—"
        assert chu_giay({"srt_start": "", "srt_end": "", "duration": ""}) == "—"

    def test_chi_co_do_dai_thi_noi_do_dai(self):
        """`4-canh.json` của vài lượt cũ chỉ có `duration` — vẫn nói được."""
        from ui_qt.bang_canh_auto import chu_giay

        assert chu_giay({"duration": 4.0}) == "4s"


# ── Bảng cảnh dùng chung ─────────────────────────────────────────────────────

@pytest.fixture(scope="module")
def qt_app():
    os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
    from PyQt5.QtWidgets import QApplication

    yield QApplication.instance() or QApplication([])


def _canh(n=3, moc=True):
    ra = []
    for i in range(1, n + 1):
        c = {"scene_id": i, "img_prompt": "anh {0}".format(i),
             "video_prompt": "clip {0}".format(i),
             "srt_text": "cau {0}".format(i)}
        if moc:
            c.update({"srt_start": "00:00:{0:02d},000".format((i - 1) * 8),
                      "srt_end": "00:00:{0:02d},000".format(i * 8),
                      "duration": 8.0})
        ra.append(c)
    return ra


class TestBangDungChung:
    def test_cot_giay_hien_tren_bang(self, qt_app):
        from ui_qt.bang_canh_auto import COT_GIAY, BangCanhWidget

        w = BangCanhWidget(canh=_canh())
        assert w._bang.item(0, COT_GIAY).text() == "00:00–00:08 · 8s"
        assert w._bang.item(2, COT_GIAY).text() == "00:16–00:24 · 8s"
        # Rê chuột là thấy nguyên văn con số trong file.
        assert "00:00:16,000" in w._bang.item(2, COT_GIAY).toolTip()

    def test_canh_khong_co_moc_thi_hien_gach_ngang(self, qt_app):
        from ui_qt.bang_canh_auto import COT_GIAY, BangCanhWidget

        w = BangCanhWidget(canh=_canh(moc=False))
        assert w._bang.item(0, COT_GIAY).text() == "—"

    def test_tab_tu_dong_khong_co_o_tick(self, qt_app):
        """Bên Tự động, SỬA CHỮ là lệnh — thêm ô tick là thêm một đường vào."""
        from ui_qt.bang_canh_auto import COT_CHON, BangCanhWidget

        w = BangCanhWidget(canh=_canh())
        assert w._bang.isColumnHidden(COT_CHON)
        assert BangCanhWidget(canh=_canh(), chon_nhieu=True) \
            ._bang.isColumnHidden(COT_CHON) is False

    def test_tick_roi_bam_thi_giao_dung_may_canh_do(self, qt_app):
        from PyQt5.QtCore import Qt

        from ui_qt.bang_canh_auto import COT_CHON, KIEU_CA_HAI, BangCanhWidget

        goi = []
        w = BangCanhWidget(canh=_canh(4), chon_nhieu=True,
                           tao_lai=lambda ds, kieu: goi.append((ds, kieu)),
                           gia_cua=lambda ds, kieu: "giá {0} cảnh".format(len(ds)))
        assert not w._nut_chon.isEnabled()
        w._bang.item(1, COT_CHON).setCheckState(Qt.Checked)
        w._bang.item(3, COT_CHON).setCheckState(Qt.Checked)
        assert w.da_tick() == [2, 4]
        assert w._nut_chon.isEnabled()
        assert "(2)" in w._nut_chon.text()
        assert w._nhan_gia.text() == "giá 2 cảnh"      # giá nói TRƯỚC khi bấm
        w._giao_chon()
        assert goi == [([2, 4], KIEU_CA_HAI)]

    def test_nap_bang_khac_thi_ve_lai_tu_dau(self, qt_app):
        from ui_qt.bang_canh_auto import COT_LOI_ANH, BangCanhWidget

        w = BangCanhWidget(canh=_canh(2))
        w.dat_canh(_canh(5))
        assert w._bang.rowCount() == 5
        assert w._bang.item(4, COT_LOI_ANH).text() == "anh 5"
        assert w._da_sua() == {}       # đổ bảng không phải người dùng gõ

    def test_luu_xong_thi_het_dau_da_sua(self, qt_app):
        from ui_qt.bang_canh_auto import COT_LOI_ANH, BangCanhWidget

        w = BangCanhWidget(canh=_canh(3))
        w._bang.setCurrentCell(0, COT_LOI_ANH)
        w._sua_anh.setPlainText("anh 1 moi")
        assert w.loi_nhac_hien()[1][0] == "anh 1 moi"
        assert w._da_sua() == {1: ("anh 1 moi", None)}
        w.chot_goc()
        assert w._da_sua() == {}


# ── Cả tab: mở tệp cũ, tick rồi tạo lại, lưu ─────────────────────────────────

class _AppGia:
    """Vừa đủ cho `TrangPromptVisuals` dựng lên — không mạng, không ví thật."""

    def __init__(self, tmp):
        self.base_dir = GOC
        self.client = object()          # coi như đã đăng nhập
        self.prices = {}
        self.config = SimpleNamespace(api_key="", base_url="")
        self.bao = []
        self.loi = []
        self.me = []
        self._tmp = str(tmp)

    def default_output_dir(self, _ten=""):
        return self._tmp

    def show_message(self, tieu_de, noi_dung):
        self.bao.append((tieu_de, noi_dung))

    def show_error(self, loi):
        self.loi.append(str(loi))

    def start_batch(self, specs, folder=""):
        self.me.append((list(specs), folder))

    def run_bg(self, ham, on_ok=None, on_err=None):
        try:
            ket = ham()
        except Exception as loi:        # noqa: BLE001 — test chạy thẳng, không luồng
            if on_err:
                on_err(loi)
            return
        if on_ok:
            on_ok(ket)

    def goi_tren_luong_ve(self, ham):
        ham()

    def bao_can_khoa(self, *_a, **_k):
        return False


def _render(duong, scenes):
    """File Excel THẬT, do chính tool sinh ra — đủ cột như lúc khách chạy."""
    from pathlib import Path

    duong_run = os.path.join(GOC, "tool-catalog", "prompt.workbook", "run.py")
    spec = importlib.util.spec_from_file_location("pw_run_bang_canh", duong_run)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    mod.render_workbook(Path(str(duong)), {"scenes": scenes, "characters": []})
    return str(duong)


@pytest.fixture
def trang(qt_app, tmp_path, monkeypatch):
    from ui_qt.trang_prompt_visuals import TrangPromptVisuals

    app = _AppGia(tmp_path)
    t = TrangPromptVisuals(app)
    t._thu_muc.dat(str(tmp_path))
    # Tải ảnh lên mạng: thay bằng một link giả, không đụng Internet.
    monkeypatch.setattr("ui_qt.trang_prompt_visuals.tai_len",
                        lambda _c, d: "https://vi-du/" + os.path.basename(d))
    yield t, app
    t.deleteLater()


def _mo_file(trang_va_app, duong, monkeypatch):
    from PyQt5.QtWidgets import QFileDialog

    monkeypatch.setattr(QFileDialog, "getOpenFileName",
                        staticmethod(lambda *a, **k: (duong, "")))
    trang_va_app[0]._mo_tep_cu()


class TestMoTepCu:
    def test_mo_file_cu_thi_bang_canh_hien_du_gio_giac(self, trang, tmp_path,
                                                       monkeypatch):
        from ui_qt.bang_canh_auto import COT_GIAY, COT_LOI_ANH

        t, _app = trang
        duong = _render(str(tmp_path / "a-prompts.xlsx"), [
            {"scene_id": 1, "srt_start": "00:00:00,000",
             "srt_end": "00:00:07,000", "duration": 7,
             "img_prompt": "anh mot", "video_prompt": "clip mot"},
            {"scene_id": 2, "srt_start": "00:01:12,000",
             "srt_end": "00:01:19,000", "duration": 7,
             "img_prompt": "anh hai", "video_prompt": "clip hai"},
        ])
        _mo_file(trang, duong, monkeypatch)
        assert t._bang_canh._bang.rowCount() == 2
        assert t._bang_canh._bang.item(1, COT_GIAY).text() == "01:12–01:19 · 7s"
        assert t._bang_canh._bang.item(1, COT_LOI_ANH).text() == "anh hai"
        assert t._nut_luu.isEnabled()

    def test_file_khong_phai_bang_canh_thi_noi_thang_thieu_cot(
            self, trang, tmp_path, monkeypatch):
        from openpyxl import Workbook

        t, app = trang
        duong = str(tmp_path / "linh-tinh.xlsx")
        wb = Workbook()
        ws = wb.active
        ws.title = "scenes"
        ws.append(["scene_id", "ghi_chu"])
        ws.append([1, "không phải bảng cảnh"])
        wb.save(duong)
        _mo_file(trang, duong, monkeypatch)
        assert app.bao, "mở file sai mà không nói gì"
        tieu_de, noi_dung = app.bao[-1]
        assert "img_prompt" in noi_dung and "video_prompt" in noi_dung
        assert t._bang_canh._bang.rowCount() == 0


class TestTaoLaiCanhDaChon:
    def _mo(self, trang, tmp_path, monkeypatch, so_canh=5):
        duong = _render(str(tmp_path / "b-prompts.xlsx"), [
            {"scene_id": i, "srt_start": "00:00:{0:02d},000".format((i - 1) * 8),
             "srt_end": "00:00:{0:02d},000".format(i * 8), "duration": 8,
             "img_prompt": "anh {0}".format(i),
             "video_prompt": "clip {0}".format(i)}
            for i in range(1, so_canh + 1)])
        _mo_file(trang, duong, monkeypatch)
        return duong

    def test_tick_hai_canh_thi_gui_dung_hai_canh_do(self, trang, tmp_path,
                                                    monkeypatch):
        from PyQt5.QtCore import Qt

        from ui_qt.bang_canh_auto import COT_CHON

        t, app = trang
        self._mo(trang, tmp_path, monkeypatch)
        bang = t._bang_canh._bang
        bang.item(2, COT_CHON).setCheckState(Qt.Checked)   # cảnh 3
        bang.item(4, COT_CHON).setCheckState(Qt.Checked)   # cảnh 5
        # Giá phải hiện TRƯỚC khi bấm.
        assert "tạm giữ" in t._bang_canh._nhan_gia.text()
        assert "3, 5" in t._bang_canh._nhan_gia.text()
        t._bang_canh._giao_chon()
        assert len(app.me) == 1
        specs, _folder = app.me[0]
        assert [s.index for s in specs] == [3, 5]
        assert [s.content for s in specs] == ["anh 3", "anh 5"]

    def test_chi_lam_lai_anh_thi_khong_tu_gui_clip(self, trang, tmp_path,
                                                   monkeypatch):
        """Mỗi clip là một lần trừ tiền cho thứ khách vừa bảo đừng làm."""
        from PyQt5.QtCore import Qt

        from core.jobs import STATUS_DONE
        from ui_qt.bang_canh_auto import COT_CHON, KIEU_ANH

        t, app = trang
        self._mo(trang, tmp_path, monkeypatch)
        t._bang_canh._bang.item(0, COT_CHON).setCheckState(Qt.Checked)
        t._bang_canh._o_kieu.setCurrentIndex(
            t._bang_canh._o_kieu.findData(KIEU_ANH))
        t._bang_canh._giao_chon()
        spec = app.me[0][0][0]
        anh = str(tmp_path / "1.png")
        with open(anh, "wb") as tep:
            tep.write(b"x")
        t.nhan_su_kien("job", SimpleNamespace(
            spec=spec, status=STATUS_DONE, files=[anh], urls=[], message=""))
        t.cuoi_nhip()
        assert len(app.me) == 1, "chọn 'chỉ làm lại ảnh' mà vẫn gửi clip"

    def test_ca_hai_thi_anh_xong_moi_noi_sang_clip(self, trang, tmp_path,
                                                   monkeypatch):
        """Ảnh xong → clip của ĐÚNG cảnh ấy chạy tiếp, dùng link máy chủ."""
        from PyQt5.QtCore import Qt

        from core.jobs import STATUS_DONE
        from core.pricing import KIND_VIDEO
        from ui_qt.bang_canh_auto import COT_CHON

        t, app = trang
        self._mo(trang, tmp_path, monkeypatch)
        t._bang_canh._bang.item(2, COT_CHON).setCheckState(Qt.Checked)  # cảnh 3
        t._bang_canh._giao_chon()
        anh = str(tmp_path / "3.png")
        with open(anh, "wb") as tep:
            tep.write(b"x")
        t.nhan_su_kien("job", SimpleNamespace(
            spec=app.me[0][0][0], status=STATUS_DONE, files=[anh],
            urls=["https://cdn.shopapi.vn/3.png"], message=""))
        t.cuoi_nhip()
        assert len(app.me) == 2, t._log.toPlainText()
        clip = app.me[1][0][0]
        assert clip.kind == KIND_VIDEO
        assert clip.index == 3 and clip.content == "clip 3"

    def test_anh_xong_thi_ghi_duong_vao_excel(self, trang, tmp_path,
                                              monkeypatch):
        from PyQt5.QtCore import Qt

        from core.jobs import STATUS_DONE
        from core.prompt_visuals import canh_de_xem
        from ui_qt.bang_canh_auto import COT_CHON, KIEU_ANH

        t, app = trang
        duong = self._mo(trang, tmp_path, monkeypatch)
        t._bang_canh._bang.item(1, COT_CHON).setCheckState(Qt.Checked)  # cảnh 2
        t._bang_canh._o_kieu.setCurrentIndex(
            t._bang_canh._o_kieu.findData(KIEU_ANH))
        t._bang_canh._giao_chon()
        anh = str(tmp_path / "2.png")
        with open(anh, "wb") as tep:
            tep.write(b"x")
        t.nhan_su_kien("job", SimpleNamespace(
            spec=app.me[0][0][0], status=STATUS_DONE, files=[anh], urls=[],
            message=""))
        from openpyxl import load_workbook

        wb = load_workbook(duong, read_only=True, data_only=True)
        try:
            hang = [list(r) for r in wb["scenes"].iter_rows(values_only=True)]
        finally:
            wb.close()
        canh = {c["scene_id"]: c for c in canh_de_xem(hang)}
        assert canh[2]["img_path"] == anh
        assert canh[1]["img_path"] == ""      # cảnh không chọn: không đụng tới

    def test_chua_dang_nhap_thi_noi_that_chu_khong_gui(self, trang, tmp_path,
                                                       monkeypatch):
        from PyQt5.QtCore import Qt

        from ui_qt.bang_canh_auto import COT_CHON

        t, app = trang
        self._mo(trang, tmp_path, monkeypatch)
        app.client = None
        t._bang_canh._bang.item(0, COT_CHON).setCheckState(Qt.Checked)
        t._bang_canh._giao_chon()
        assert app.me == []
        assert app.bao[-1][0] == "Chưa đăng nhập"


class TestLuuChinhSua:
    def test_sua_tren_bang_roi_luu_thi_ghi_dung_cot_theo_ten(
            self, trang, tmp_path, monkeypatch):
        from core.prompt_visuals import canh_de_xem
        from ui_qt.bang_canh_auto import COT_LOI_ANH

        t, app = trang
        duong = _render(str(tmp_path / "c-prompts.xlsx"), [
            {"scene_id": 1, "srt_start": "00:00:00,000",
             "srt_end": "00:00:08,000", "duration": 8,
             "img_prompt": "anh mot", "video_prompt": "clip mot"},
            {"scene_id": 2, "srt_start": "00:00:08,000",
             "srt_end": "00:00:16,000", "duration": 8,
             "img_prompt": "anh hai", "video_prompt": "clip hai"},
        ])
        _mo_file(trang, duong, monkeypatch)
        t._bang_canh._bang.setCurrentCell(1, COT_LOI_ANH)
        t._bang_canh._sua_anh.setPlainText("anh hai da sua")
        t._luu_chinh_sua()
        assert not app.loi, app.loi

        from openpyxl import load_workbook

        wb = load_workbook(duong, read_only=True, data_only=True)
        try:
            hang = [list(r) for r in wb["scenes"].iter_rows(values_only=True)]
        finally:
            wb.close()
        canh = {c["scene_id"]: c for c in canh_de_xem(hang)}
        assert canh[2]["img_prompt"] == "anh hai da sua"
        assert canh[2]["video_prompt"] == "clip hai"   # không đụng cột khác
        assert canh[2]["srt_start"] == "00:00:08,000"  # mốc giờ giữ nguyên
        assert canh[1]["img_prompt"] == "anh mot"
        # Lưu xong thì hết dấu "đã sửa".
        assert t._bang_canh._da_sua() == {}


def test_trang_khong_tran_mep_760px(trang, qt_app):
    """CLAUDE.md: trang không co xuống 760px là có phần bị đẩy ra ngoài mép."""
    t, _app = trang
    t.resize(760, 900)
    qt_app.processEvents()
    assert t.minimumSizeHint().width() <= 760, t.minimumSizeHint().width()
